# build_charter_files.py — สร้างไฟล์ Word / PDF / สารบัญ ของรัฐธรรมนูญ ๒๕๒๕ จากตัวบทใน website_new.html
#
#   python tools/build_charter_files.py
#
# ตัวบทอ่านจาก website_new.html (สร้างโดย tools/build_site.py) ซึ่งรวมการแก้ไขเพิ่มเติมทุกรอบแล้ว
# ไม่ได้ใช้ files/processed_constitution.txt แบบ Genarate/generate_*.py ที่ตกรุ่นไปแล้ว
# ผลลัพธ์:
#   files/รัฐธรรมนูญ_๒๕๒๕_แก้ไขเพิ่มเติม.docx   (รูปแบบเดียวกับ Genarate/generate_word.py + ตราครุฑหน้าแรก)
#   files/สารบัญรัฐธรรมนูญ.docx
#   constitution_2525.html                        (หน้าพิมพ์แบบราชกิจจานุเบกษา ใช้พิมพ์ PDF)
#   files/รัฐธรรมนูญ_๒๕๒๕_แก้ไขเพิ่มเติม.pdf      (พิมพ์จาก constitution_2525.html ด้วย Chromium)
# แล้วอ่านทั้งสามไฟล์กลับมาเทียบกับตัวบทในเว็บทีละมาตรา หยุดทันทีถ้าไม่ตรง

import html
import os
import re
import subprocess
import sys

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
SRC = os.path.join(ROOT, 'website_new.html')
BASE = 'รัฐธรรมนูญ_๒๕๒๕_แก้ไขเพิ่มเติม'
WORD_OUT = os.path.join(ROOT, 'files', BASE + '.docx')
TOC_OUT = os.path.join(ROOT, 'files', 'สารบัญรัฐธรรมนูญ.docx')
HTML_OUT = os.path.join(ROOT, 'constitution_2525.html')
PDF_OUT = os.path.join(ROOT, 'files', BASE + '.pdf')
GARUDA_PNG = os.path.join(ROOT, 'images', 'Thai_Garuda_emblem.png')
GARUDA_SVG = 'images/Thai_Garuda_emblem.svg'      # relative to constitution_2525.html
THAI_FONT = 'TH Sarabun New'
FULL_TITLE = 'รัฐธรรมนูญแห่งราชอาณาจักรไทย พุทธศักราช ๒๕๒๕'

try:
    sys.stdout.reconfigure(encoding='utf-8')
except AttributeError:
    pass


def thai(n):
    return str(n).translate(str.maketrans('0123456789', '๐๑๒๓๔๕๖๗๘๙'))


def text_of(fragment):
    """Visible text of an HTML fragment: tooltips dropped, <br> kept as newlines."""
    fragment = re.sub(r'<span class="tooltiptext">.*?</span>', '', fragment, flags=re.S)
    fragment = re.sub(r'<br\s*/?>', '\n', fragment)
    fragment = html.unescape(re.sub(r'<[^>]+>', '', fragment))
    return '\n'.join(re.sub(r'[ \t]+', ' ', line).strip() for line in fragment.split('\n')).strip()


def split_lines(lines):
    """Turn loose lines into ('list', num, text) / ('para', text) items."""
    items = []
    for line in lines:
        line = line.strip()
        if not line:
            continue
        m = re.match(r'^(\([๐-๙]+\))\s*(.*)', line, re.S)
        items.append(('list', m.group(1), m.group(2)) if m else ('para', line))
    return items


# ==================== Read the charter from website_new.html ====================

src = open(SRC, encoding='utf-8').read()
pre = src[src.index('<section class="nw-preamble"'):]
pre = pre[:pre.index('</section>')]
title = text_of(re.search(r'<h2>(.*?)</h2>', pre, re.S).group(1))
king_lines = text_of(re.search(r'<p class="nw-preamble-sub">(.*?)</p>', pre, re.S).group(1)).split('\n')
preamble = [text_of(p) for p in re.findall(r'<p>(.*?)</p>', pre, re.S)]

chapters = []
for sec in re.finditer(r'<section class="nw-chapter" id="ch_\d+">(.*?)</section>', src, re.S):
    body = sec.group(1)
    head = re.search(r'<div class="nw-ch-no">(.*?)</div><div><h2>(.*?)</h2>', body, re.S)
    ch = {'no': text_of(head.group(1)), 'name': text_of(head.group(2)), 'blocks': []}
    body = body[body.index('</header>') + 9:]
    for m in re.finditer(r'<h3 class="nw-part">(.*?)</h3>'
                         r'|<article class="nw-art[^"]*" id="art_(\d+)"[^>]*>(.*?)</article>'
                         r'|<div class="nw-clause">(.*?)</div>'
                         r'|<div class="nw-signature">(.*?)</div>', body, re.S):
        if m.group(1) is not None:
            ch['blocks'].append(('part', text_of(m.group(1)).split('\n')))
        elif m.group(2) is not None:
            art = m.group(3)
            art = art[art.index('<div class="nw-art-body">'):]
            art = re.sub(r'<div class="nw-compare".*', '', art, flags=re.S)   # the 2540 comparison is not charter text
            items = []
            for p in re.finditer(r'<p class="(article-text|sub-para|list-item)">(.*?)</p>', art, re.S):
                if p.group(1) == 'list-item':
                    num = re.match(r'\s*<span class="list-num">(.*?)</span>\s*', p.group(2), re.S)
                    items.append(('list', text_of(num.group(1)), text_of(p.group(2)[num.end():])))
                else:
                    items.append(('para', text_of(p.group(2))))
            assert items and items[0][0] == 'para', 'article %s does not open with its text' % m.group(2)
            ch['blocks'].append(('art', int(m.group(2)), items))
        elif m.group(4) is not None:
            # a clause the site shows after an article belongs to that article (e.g. ม.๒๙๕)
            assert ch['blocks'] and ch['blocks'][-1][0] == 'art', 'clause with no article before it'
            ch['blocks'][-1][2].extend(split_lines(text_of(m.group(4)).split('\n')))
        else:
            ch['blocks'].append(('sign', text_of(m.group(5)).split('\n')))
    chapters.append(ch)

arts = [b for ch in chapters for b in ch['blocks'] if b[0] == 'art']
nums = [b[1] for b in arts]
assert nums == list(range(1, 1161)), 'articles are not ๑–๑๑๖๐ in order'
assert len(chapters) == 38, len(chapters)
print('อ่านจาก website_new.html: %d หมวด, %d มาตรา' % (len(chapters), len(arts)))


def chapter_lines(ch):
    return [x for x in (ch['no'], ch['name']) if x]


def chapter_label(ch):
    return ' '.join(chapter_lines(ch))


def chapter_range(ch):
    n = [b[1] for b in ch['blocks'] if b[0] == 'art']
    return (min(n), max(n)) if n else None


# ==================== Word ====================
#
# Every paragraph uses a named style and runs carry no formatting except bold on "มาตรา N", so the
# file stays clean when edited in Word. Thai text is a complex script: Word reads its font, size and
# bold from w:cs / w:szCs / w:bCs, so each of those is set next to its Latin twin, in schema order.

def _el(tag, **attrs):
    el = OxmlElement(tag)
    for k, v in attrs.items():
        el.set(qn('w:' + k), v)
    return el


def cs_twins(rPr, bold=None):
    """Add w:bCs after w:b and w:szCs after w:sz (Word ignores them when they sit out of order)."""
    sz = rPr.find(qn('w:sz'))
    if sz is not None and rPr.find(qn('w:szCs')) is None:
        sz.addnext(_el('w:szCs', val=sz.get(qn('w:val'))))
    b = rPr.find(qn('w:b'))
    if b is not None and rPr.find(qn('w:bCs')) is None:
        b.addnext(_el('w:bCs', **({'val': b.get(qn('w:val'))} if b.get(qn('w:val')) else {})))


def thai_fonts(rPr):
    """Explicit TH Sarabun New for every script, with the theme-font attributes removed."""
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = _el('w:rFonts')
        rPr.insert(0, rFonts)
    for k in list(rFonts.attrib):
        del rFonts.attrib[k]
    for k in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        rFonts.set(qn('w:' + k), THAI_FONT)


def set_defaults(doc):
    styles = doc.styles.element
    defaults = styles.find(qn('w:docDefaults'))
    for child in list(defaults):
        defaults.remove(child)
    rPr = _el('w:rPr')
    rPr.append(_el('w:rFonts', ascii=THAI_FONT, hAnsi=THAI_FONT, eastAsia=THAI_FONT, cs=THAI_FONT))
    rPr.append(_el('w:noProof'))                         # no red underlines on Thai
    rPr.append(_el('w:sz', val='32'))
    rPr.append(_el('w:szCs', val='32'))
    rPr.append(_el('w:lang', val='th-TH', eastAsia='th-TH', bidi='th-TH'))
    rd = _el('w:rPrDefault')
    rd.append(rPr)
    pPr = _el('w:pPr')
    pPr.append(_el('w:spacing', after='0', line='240', lineRule='auto'))
    pd = _el('w:pPrDefault')
    pd.append(pPr)
    defaults.append(rd)
    defaults.append(pd)
    normal = doc.styles['Normal']
    thai_fonts(normal.element.get_or_add_rPr())
    normal.font.size = Pt(16)
    cs_twins(normal.element.rPr)
    pf = normal.paragraph_format
    pf.space_before, pf.space_after, pf.line_spacing = Pt(0), Pt(0), 1.0
    pf.widow_control = True


def add_style(doc, name, size=16, bold=False, align=None, before=0, after=0, first=None, left=None,
              keep_next=False, page_break=False, outline=None):
    st = doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    st.base_style = doc.styles['Normal']
    st.quick_style = True
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor(0, 0, 0)
    if bold:
        st.font.bold = True
    thai_fonts(st.element.get_or_add_rPr())
    cs_twins(st.element.rPr)
    pf = st.paragraph_format
    if align is not None:
        pf.alignment = align
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    if first is not None:
        pf.first_line_indent = Cm(first)
    if left is not None:
        pf.left_indent = Cm(left)
    pf.keep_with_next = keep_next
    pf.page_break_before = page_break
    if outline is not None:                              # shows in Word's navigation pane
        st.element.get_or_add_pPr().append(_el('w:outlineLvl', val=str(outline)))
    return st


def bold_run(paragraph, text):
    run = paragraph.add_run(text)
    run.font.bold = True
    cs_twins(run._element.get_or_add_rPr())
    return run


def add_lines(paragraph, lines):
    for i, line in enumerate(lines):
        run = paragraph.add_run(line)
        if i < len(lines) - 1:
            run.add_break()


def field(paragraph, code):
    for kind in ('begin', None, 'end'):
        run = paragraph.add_run()
        if kind:
            run._element.append(_el('w:fldChar', fldCharType=kind))
        else:
            instr = _el('w:instrText')
            instr.set(qn('xml:space'), 'preserve')
            instr.text = ' %s ' % code
            run._element.append(instr)


def header_line(paragraph, left, center, right_text=None, page=False, bold_center=False):
    """left <tab> centre <tab> right, on one line with a rule under it (the gazette header)."""
    paragraph.style = 'หัวกระดาษ'
    paragraph.add_run(left + '\t')
    if bold_center:
        bold_run(paragraph, center)
    else:
        paragraph.add_run(center)
    paragraph.add_run('\t' + (right_text or ''))
    if page:
        field(paragraph, 'PAGE')                        # Thai digits come from the section's pgNumType


TEXT_W = 16.0          # A4 21 cm − left 3 cm − right 2 cm


def setup_page(doc, gazette_header=True):
    section = doc.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.5), Cm(2)
    section.left_margin, section.right_margin = Cm(3), Cm(2)
    section.header_distance, section.footer_distance = Cm(1.25), Cm(1)
    hs = add_style(doc, 'หัวกระดาษ', size=12, after=0)
    hs.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W / 2), WD_TAB_ALIGNMENT.CENTER)
    hs.paragraph_format.tab_stops.add_tab_stop(Cm(TEXT_W), WD_TAB_ALIGNMENT.RIGHT)
    border = _el('w:pBdr')
    border.append(_el('w:bottom', val='single', sz='6', space='4', color='000000'))
    hs.element.get_or_add_pPr().append(border)
    section._sectPr.append(_el('w:pgNumType', fmt='thaiNumbers'))
    if gazette_header:
        section.different_first_page_header_footer = True
        header_line(section.first_page_header.paragraphs[0], GAZETTE_ISSUE, 'ราชกิจจานุเบกษา', GAZETTE_DATE,
                    bold_center=True)
        header_line(section.header.paragraphs[0], GAZETTE_ISSUE, 'ราชกิจจานุเบกษา', 'หน้า ', page=True)
    else:
        p = section.header.paragraphs[0]
        p.style = 'หัวกระดาษ'
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run(FULL_TITLE)


def new_document(gazette_header=True):
    doc = Document()
    set_defaults(doc)
    setup_page(doc, gazette_header)
    doc.core_properties.title = FULL_TITLE
    doc.core_properties.author = 'รัฐธรรมนุญจำลอง'
    lang = doc.settings.element.find(qn('w:themeFontLang'))
    if lang is not None:
        lang.set(qn('w:val'), 'th-TH')
        lang.set(qn('w:bidi'), 'th-TH')
    return doc


GAZETTE_ISSUE = 'เล่ม ๙๙ ตอนที่ ๑๔๕ ก'
GAZETTE_DATE = '๕ ธันวาคม ๒๕๒๕'
# Left-aligned on purpose: Thai distributed (and justify) letter-spaced whole lines in the user's Word,
# which has no Thai word breaking and can only break at spaces ("ก า ร ติ ช ม", "ม า ต ร า ๑ ๑").
DIST = WD_ALIGN_PARAGRAPH.LEFT

doc = new_document()
add_style(doc, 'ตราครุฑ', align=WD_ALIGN_PARAGRAPH.CENTER, before=6, after=6, keep_next=True)
add_style(doc, 'ชื่อรัฐธรรมนูญ', size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=6, keep_next=True)
add_style(doc, 'พระปรมาภิไธย', size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)
add_style(doc, 'ย่อหน้า', align=DIST, first=2.5, after=3)
add_style(doc, 'หมวด', size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=12,
          keep_next=True, page_break=True, outline=0)
add_style(doc, 'ส่วน', bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=6, keep_next=True, outline=1)
add_style(doc, 'มาตรา', align=DIST, first=2.5, before=6, after=0)
add_style(doc, 'วรรค', align=DIST, first=2.5, after=0)
add_style(doc, 'อนุมาตรา', align=DIST, left=3.5, first=-1)
add_style(doc, 'ผู้รับสนอง', align=WD_ALIGN_PARAGRAPH.CENTER, left=8, keep_next=True)


def para(style, text=None):
    p = doc.add_paragraph(style=style)
    if text is not None:
        p.add_run(text)
    return p


para('ตราครุฑ').add_run().add_picture(GARUDA_PNG, width=Cm(3))
para('ชื่อรัฐธรรมนูญ', title)
add_lines(para('พระปรมาภิไธย'), king_lines)
for text in preamble:
    para('ย่อหน้า', text)

for ch in chapters:
    add_lines(para('หมวด'), chapter_lines(ch))           # every chapter opens a new page
    for b in ch['blocks']:
        if b[0] == 'part':
            add_lines(para('ส่วน'), b[1])
        elif b[0] == 'art':
            for k, item in enumerate(b[2]):
                if item[0] == 'list':
                    para('อนุมาตรา', '%s\t%s' % (item[1], item[2]))
                elif k == 0:
                    p = para('มาตรา')
                    bold_run(p, 'มาตรา %s' % thai(b[1]))
                    p.add_run('  ' + item[1])
                else:
                    para('วรรค', item[1])
        else:
            first = para('ผู้รับสนอง', b[1][0])
            first.paragraph_format.space_before = Pt(36)
            for line in b[1][1:]:
                bold_run(para('ผู้รับสนอง'), line)

doc.save(WORD_OUT)
print('Word: %s (%s bytes)' % (os.path.relpath(WORD_OUT, ROOT), format(os.path.getsize(WORD_OUT), ',')))

# ==================== Table of contents ====================

toc = new_document(gazette_header=False)
add_style(toc, 'สารบัญ', size=26, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, before=12, after=0)
add_style(toc, 'สารบัญ-ชื่อ', size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, after=18)
add_style(toc, 'ตาราง', size=15)
add_style(toc, 'ตาราง-กลาง', size=15, align=WD_ALIGN_PARAGRAPH.CENTER)
add_style(toc, 'สรุป', bold=True, align=WD_ALIGN_PARAGRAPH.RIGHT, before=12)
toc.add_paragraph('สารบัญ', style='สารบัญ')
toc.add_paragraph(FULL_TITLE, style='สารบัญ-ชื่อ')
table = toc.add_table(rows=1, cols=3)
table.style = 'Table Grid'
table.autofit = False
table.alignment = WD_TABLE_ALIGNMENT.CENTER
widths = (Cm(9.6), Cm(3.2), Cm(3.2))
for col, w in zip(table.columns, widths):
    col.width = w      # the grid widths; LibreOffice ignores per-cell widths alone


def shade(cell, fill):
    cell._element.get_or_add_tcPr().append(_el('w:shd', val='clear', color='auto', fill=fill))


def cell_text(cell, w, text, style, bold=False, white=False):
    cell.width = w
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    p.style = style
    run = bold_run(p, text) if bold else p.add_run(text)
    if white:
        run.font.color.rgb = RGBColor(255, 255, 255)


head = table.rows[0]
trPr = head._tr.get_or_add_trPr()
trPr.append(_el('w:tblHeader'))                          # header row repeats on each page
for cell, w, label in zip(head.cells, widths, ('หมวด', 'มาตราแรก', 'มาตราท้าย')):
    shade(cell, '2C3E50')
    cell_text(cell, w, label, 'ตาราง-กลาง', bold=True, white=True)
toc_rows = []
for i, ch in enumerate(chapters):
    first, last = chapter_range(ch)
    toc_rows.append((chapter_label(ch), first, last))
    row = table.add_row()
    row._tr.get_or_add_trPr().append(_el('w:cantSplit'))
    for cell, w, value, style in zip(row.cells, widths, (chapter_label(ch), thai(first), thai(last)),
                                     ('ตาราง', 'ตาราง-กลาง', 'ตาราง-กลาง')):
        shade(cell, 'EBF5FB' if i % 2 == 0 else 'FFFFFF')
        cell_text(cell, w, value, style)
toc.add_paragraph('รวมทั้งสิ้น %s หมวด  %s มาตรา' % (thai(len(chapters)), thai(len(arts))), style='สรุป')
toc.save(TOC_OUT)
print('สารบัญ: %s' % os.path.relpath(TOC_OUT, ROOT))

# ==================== Print page (constitution_2525.html) ====================

esc = html.escape


def items_html(items, art_no):
    out = []
    for k, item in enumerate(items):
        if item[0] == 'list':
            out.append('<p class="list-item"><span class="list-num">%s</span>%s</p>' % (esc(item[1]), esc(item[2])))
        elif k == 0:
            out.append('<p class="article"><span class="bold">มาตรา %s</span> %s</p>' % (thai(art_no), esc(item[1])))
        else:
            out.append('<p class="sub-paragraph">%s</p>' % esc(item[1]))
    return '\n'.join(out)


body = ['''<div class="first-page-header">
  <table class="gazette-meta"><tr>
    <td style="text-align:left;width:35%">เล่ม ๙๙ ตอนที่ ๑๔๕ ก</td>
    <td style="text-align:center;width:30%;font-weight:bold">ราชกิจจานุเบกษา</td>
    <td style="text-align:right;width:35%">''' + GAZETTE_DATE + '''</td>
  </tr></table>
  <hr class="double-line">
  <div class="garuda-container"><img src="''' + GARUDA_SVG + '''" class="garuda-img" alt=""></div>
  <div class="main-title">''' + esc(title) + '''</div>
  <div class="king-title">''' + '<br>'.join(map(esc, king_lines)) + '''</div>
</div>''', '<div class="preamble-container">']
body += ['<p class="preamble">%s</p>' % esc(p) for p in preamble]
body.append('</div>')
for ch in chapters:
    body.append('<div class="chapter-title">%s</div>' % '<br>'.join(map(esc, chapter_lines(ch))))
    for b in ch['blocks']:
        if b[0] == 'part':
            body.append('<div class="part-title">%s</div>' % '<br>'.join(map(esc, b[1])))
        elif b[0] == 'art':
            body.append(items_html(b[2], b[1]))
        else:
            body.append('<div class="signature-container">%s</div>' % '<br>'.join(map(esc, b[1])))

page = '''<!DOCTYPE html>
<html lang="th">
<head>
<meta charset="utf-8">
<title>''' + FULL_TITLE + '''</title>
<!-- สร้างโดย tools/build_charter_files.py จาก website_new.html — อย่าแก้ไฟล์นี้ด้วยมือ -->
<style>
  @import url('https://fonts.googleapis.com/css2?family=Sarabun:wght@400;700&display=swap');
  @page {
    size: A4;
    margin: 3cm 2.5cm 2.5cm 2.5cm;
    @top-left { content: "เล่ม ๙๙ ตอนที่ ๑๔๕ ก"; font-family: 'Sarabun', sans-serif; font-size: 10pt; border-bottom: 0.5px solid #000; padding-bottom: 5px; }
    @top-center { content: "ราชกิจจานุเบกษา"; font-family: 'Sarabun', sans-serif; font-size: 10pt; border-bottom: 0.5px solid #000; padding-bottom: 5px; }
    @top-right { content: "หน้า " counter(page, thai); font-family: 'Sarabun', sans-serif; font-size: 10pt; border-bottom: 0.5px solid #000; padding-bottom: 5px; }
  }
  @page :first {
    margin-top: 1.5cm;
    @top-left { content: none; } @top-center { content: none; } @top-right { content: none; }
  }
  body { font-family: 'Sarabun', sans-serif; font-size: 15pt; line-height: 1.6; color: #000; margin: 0; padding: 0; }
  /* Chromium justifies Thai only by stretching the spaces between words, which leaves wide gaps, so align left */
  p { text-align: left; margin: 0.3em 0; }
  /* Chromium paints the stacked Thai marks of a line pushed to the next page at the foot of this one;
     clipping each paragraph fragment to its own box removes those stray marks */
  p, .chapter-title, .part-title, .signature-container { overflow: hidden; }
  p.preamble, p.normal { text-indent: 1.5cm; }
  p.article { text-indent: 1.5cm; margin: 0.4em 0; }
  p.sub-paragraph { text-indent: 1.5cm; margin: 0.2em 0; }
  p.list-item { text-indent: 0; margin: 0.2em 0 0.2em 1.2cm; padding-left: 1cm; }   /* number sits in the padding, inside the clip */
  span.list-num { display: inline-block; width: 1cm; margin-left: -1cm; }
  .bold { font-weight: bold; }
  .chapter-title { font-size: 16pt; font-weight: bold; text-align: center; margin: 0 0 0.8em; break-before: page; break-after: avoid; line-height: 1.4; }
  .part-title { font-size: 15pt; font-weight: bold; text-align: center; margin: 1.2em 0 0.6em; page-break-after: avoid; line-height: 1.4; }
  .first-page-header { text-align: center; margin-bottom: 1.5em; }
  .gazette-meta { width: 100%; border-collapse: collapse; margin-bottom: 5px; font-size: 11pt; }
  .gazette-meta td { padding: 0; }
  .double-line { border: none; border-top: 1px solid #000; border-bottom: 3px double #000; height: 4px; margin: 0 0 1.5em; }
  .garuda-container { margin: 1em auto; text-align: center; }
  .garuda-img { width: 100px; height: auto; }
  .main-title { font-size: 20pt; font-weight: bold; margin: 0.5em 0; line-height: 1.3; }
  .king-title { font-size: 16pt; font-weight: bold; line-height: 1.7; margin-bottom: 1em; }
  .preamble-container { margin-bottom: 2em; }
  .signature-container { margin-top: 2.5em; margin-left: 55%; text-align: left; line-height: 1.7; page-break-inside: avoid; font-weight: bold; }
</style>
</head>
<body>
''' + '\n'.join(body) + '''
</body>
</html>
'''
open(HTML_OUT, 'w', encoding='utf-8').write(page)
print('หน้าพิมพ์: %s' % os.path.relpath(HTML_OUT, ROOT))

# ==================== PDF (Chromium through Playwright) ====================

node = r'''
const { chromium } = require('playwright');
(async () => {
  const p = process.env.HTTPS_PROXY;
  const opts = process.env.CHROME ? { executablePath: process.env.CHROME } : {};
  if (p) opts.proxy = { server: p };
  const browser = await chromium.launch(opts);
  const page = await browser.newPage();
  await page.goto(process.argv[1], { waitUntil: 'networkidle', timeout: 120000 });
  await page.evaluate(() => document.fonts.ready);
  const ok = await page.evaluate(() => document.fonts.check('15pt Sarabun', 'ก'));
  if (!ok) { console.error('Sarabun did not load'); process.exit(2); }
  await page.pdf({ path: process.argv[2], format: 'A4', preferCSSPageSize: true, printBackground: true });
  await browser.close();
})().catch(e => { console.error(e); process.exit(1); });
'''
env = dict(os.environ)
env.setdefault('NODE_PATH', subprocess.run(['npm', 'root', '-g'], capture_output=True, text=True).stdout.strip())
res = subprocess.run(['node', '-e', node, 'file://' + HTML_OUT, PDF_OUT], env=env, capture_output=True, text=True)
if res.returncode:
    sys.exit('PDF failed: ' + (res.stderr or res.stdout)[-800:])
print('PDF: %s (%s bytes)' % (os.path.relpath(PDF_OUT, ROOT), format(os.path.getsize(PDF_OUT), ',')))

# ==================== Read every output back and compare with the site ====================

squash = lambda s: re.sub(r'\s+', '', s).replace('ํา', 'ำ')
expected = squash(''.join([title] + king_lines + preamble + [
    ''.join(chapter_lines(ch)) + ''.join(
        ''.join(b[1]) if b[0] in ('part', 'sign') else
        'มาตรา' + thai(b[1]) + ''.join(''.join(it[1:]) for it in b[2])
        for b in ch['blocks'])
    for ch in chapters]))

got_word = squash(''.join(p.text for p in Document(WORD_OUT).paragraphs))
assert got_word == expected, 'Word text differs from website_new.html'
gazette_head = lambda s: re.sub(r'เล่ม๙๙ตอนที่๑๔๕กราชกิจจานุเบกษา(หน้า[๐-๙]+|' + squash(GAZETTE_DATE) + ')', '', s)
got_html = gazette_head(squash(text_of(page[page.index('<body>'):])))
assert got_html == expected, 'print page text differs from website_new.html'
import fitz  # PyMuPDF
pdf = fitz.open(PDF_OUT)
got_pdf = gazette_head(squash(''.join(pg.get_text() for pg in pdf)))
if got_pdf != expected:
    i = next((k for k, (a, b) in enumerate(zip(got_pdf, expected)) if a != b), min(len(got_pdf), len(expected)))
    sys.exit('PDF text differs at %d: pdf …%s… / site …%s…' % (i, got_pdf[i - 40:i + 40], expected[i - 40:i + 40]))
toc_doc = Document(TOC_OUT)
assert len(toc_doc.tables[0].rows) == len(chapters) + 1
print('ตรวจแล้ว: Word, หน้าพิมพ์ และ PDF (%d หน้า) ตรงกับตัวบทในเว็บทุกตัวอักษร (%s ตัว)'
      % (pdf.page_count, format(len(expected), ',')))
